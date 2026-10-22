"""输出 Excel 明细 + HTML 可视化报告。"""
from __future__ import annotations

import html
import json
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PCT_COLS = ("率", "占比", "价格分位", "比中位价")


def is_pct(col: str) -> bool:
    return any(k in col for k in PCT_COLS) and "倍率" not in col


def is_yen(col: str) -> bool:
    col = col.replace("评价", "")
    return ("价" in col or col == "price") and "变化" not in col and not is_pct(col)


def is_int(col: str) -> bool:
    return any(k in col for k in ("排名", "rank", "名次", "评价数", "评价总数", "评价增长", "销量", "review_count", "商品数"))


def _fmt(col: str, v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "-"
    if isinstance(v, (bool,)):
        return "✓" if v else ""
    if is_pct(col) and isinstance(v, (int, float)):
        return f"{v:+.0%}" if col == "比中位价" else f"{v:.0%}"
    if is_yen(col) and isinstance(v, (int, float)):
        return f"¥{v:,.0f}"
    if is_int(col) and isinstance(v, (int, float)):
        return f"{v:,.0f}"
    if isinstance(v, float):
        return f"{v:,.2f}" if abs(v) < 10 else f"{v:,.0f}"
    if isinstance(v, int):
        return f"{v:,}"
    return str(v)


# ---------------- Excel ----------------
def write_excel(result: dict, raw: pd.DataFrame, path: Path):
    sheets: dict[str, pd.DataFrame] = {"概览": pd.DataFrame(result["overview"])}

    def stack(key):
        frames = [c[key].assign(品类=c["name"]) for c in result["categories"]
                  if isinstance(c[key], pd.DataFrame) and not c[key].empty]
        if not frames:
            return pd.DataFrame()
        d = pd.concat(frames, ignore_index=True)
        return d[["品类"] + [x for x in d.columns if x != "品类"]]

    sheets["我的商品诊断"] = stack("mine")
    sheets["关键词排名"] = stack("keywords")
    sheets["价格带"] = stack("bands")
    sheets["竞争店铺"] = stack("shops")
    sheets["爆款TOP20"] = stack("bestsellers")
    sheets["标题热词"] = stack("titles")
    if result.get("trend") is not None and not result["trend"].empty:
        sheets["趋势变化"] = result["trend"]
    w = stack("watch")
    if not w.empty:
        sheets["盯防店铺商品"] = w
    sheets["原始数据"] = raw

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, d in sheets.items():
            if d.empty:
                continue
            d = d.copy()
            for c in d.columns:  # Excel 不支持 category 类型
                if str(d[c].dtype) == "category":
                    d[c] = d[c].astype(str)
            d.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="BF0000")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(d.columns, 1):
                letter = get_column_letter(i)
                width = min(max(len(str(col)) * 2, d[col].astype(str).str.len().quantile(.9) * 1.1
                                if len(d) else 8, 8), 60)
                ws.column_dimensions[letter].width = width
                pct = is_pct(str(col))
                yen = is_yen(str(col))
                for cell in ws[letter][1:]:
                    if str(col) == "诊断建议":
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                    elif pct:
                        cell.number_format = "0%"
                    elif yen:
                        cell.number_format = "¥#,##0"
                    if str(col).endswith("url") and isinstance(cell.value, str) and cell.value.startswith("http"):
                        cell.hyperlink = cell.value
                        cell.font = Font(color="0563C1", underline="single")


# ---------------- HTML ----------------
def _table(d: pd.DataFrame, cols=None, link_col="item_url", title_col=None, max_rows=50, own_col=None):
    if d is None or d.empty:
        return "<p class='muted'>暂无数据</p>"
    cols = cols or [c for c in d.columns if not str(c).endswith("url")]
    out = ["<div class='tw'><table><thead><tr>"]
    out += [f"<th>{html.escape(str(c))}</th>" for c in cols]
    out.append("</tr></thead><tbody>")
    for _, r in d.head(max_rows).iterrows():
        own = own_col and bool(r.get(own_col))
        out.append(f"<tr class='{'own' if own else ''}'>")
        for c in cols:
            txt = html.escape(_fmt(str(c), r[c])).replace("\n", "<br>")
            if c == title_col and link_col in r and isinstance(r[link_col], str) and r[link_col]:
                txt = f"<a href='{html.escape(r[link_col])}' target='_blank' rel='noopener'>{txt}</a>"
            cls = "name" if c in (title_col, "诊断建议") else ("txt" if isinstance(r[c], str) else "")
            out.append(f"<td class='{cls}'>{txt}</td>")
        out.append("</tr>")
    out.append("</tbody></table></div>")
    return "".join(out)


def write_html(result: dict, path: Path, cfg: dict):
    shops = "、".join(f"{s.get('name', s['code'])}（{s['code']}）" for s in cfg.get("my_shops", []))
    charts, sections, nav = [], [], []
    for i, c in enumerate(result["categories"]):
        ov = c["overview"]
        nav.append(f"<a href='#cat{i}'>{html.escape(c['name'])}</a>")
        kpis = [("样本商品", f"{ov['样本商品数']:,}"), ("价格中位数", _fmt("价", ov["价格中位数"])),
                ("前30名均价", _fmt("价", ov["前30名均价"])),
                ("前30名平均评价", _fmt("评价数", ov["前30名平均评价数"])),
                ("包邮率", _fmt("率", ov["包邮率"])), ("我的商品", f"{ov['我的商品数']}"),
                ("我的最好排名", "-" if pd.isna(ov["我的最好排名"]) else f"第{int(ov['我的最好排名'])}名")]
        kpi_html = "".join(f"<div class='kpi'><div class='k'>{k}</div><div class='v'>{v}</div></div>"
                           for k, v in kpis)
        ins = "".join(f"<li>{html.escape(x)}</li>" for x in c["insights"]) or "<li>暂无</li>"
        b = c["bands"]
        if not b.empty:
            charts.append({"id": f"band{i}", "labels": b["价格带"].astype(str).tolist(),
                           "items": b["商品占比"].round(3).tolist(), "reviews": b["评价占比"].round(3).tolist(),
                           "mine": b["我的商品数"].tolist()})
        mine_cols = ["店铺", "商品名", "价格", "比中位价", "评价数", "评分", "积分倍率", "包邮", "最好排名",
                     "最好排名关键词", "诊断建议"]
        band_cols = ["价格带", "商品数", "商品占比", "评价占比", "机会指数", "前30名商品数", "我的商品数", "判断"]
        shop_cols = ["名次", "店铺", "前30名商品数", "最好排名", "上榜商品数", "评价总数", "平均价格",
                     "平均评分", "包邮率", "平均积分倍率"]
        best_cols = ["商品名", "店铺", "价格", "评价数", "评分", "积分倍率", "包邮", "最好排名"]
        sections.append(f"""
<section id='cat{i}'>
  <h2>{html.escape(c['name'])}</h2>
  <div class='kpis'>{kpi_html}</div>
  <div class='card'><h3>结论与建议</h3><ul class='ins'>{ins}</ul></div>
  <div class='card'><h3>我的商品诊断</h3>{_table(c['mine'], [x for x in mine_cols if x in c['mine'].columns] if not c['mine'].empty else None, title_col='商品名')}</div>
  <div class='card'><h3>关键词排名</h3>{_table(c['keywords'])}</div>
  <div class='card'><h3>价格带分析</h3><p class='muted'>机会指数 = 评价占比 ÷ 商品占比，大于 1 表示该价格带需求比供给多</p>
    <div class='chart'><canvas id='band{i}'></canvas></div>{_table(b, band_cols if not b.empty else None)}</div>
  <div class='card'><h3>竞争店铺排行</h3>{_table(c['shops'], shop_cols if not c['shops'].empty else None, own_col='是我的店', max_rows=20)}</div>
  <div class='card'><h3>爆款 TOP20（按评价数）</h3>{_table(c['bestsellers'], best_cols if not c['bestsellers'].empty else None, title_col='商品名', own_col='是我的店')}</div>
  <div class='card'><h3>标题热词</h3>{_table(c['titles'], max_rows=30)}</div>
</section>""")

    t = result.get("trend")
    if t is not None and not t.empty:
        nav.append("<a href='#trend'>趋势</a>")
        tc = ["品类", "店铺", "商品名", "现价", "价格变化", "评价数", "评价增长", "估算月销量", "现排名", "排名变化"]
        price_moves = t[t["价格变化"] != 0].sort_values("价格变化")
        sections.append(f"""
<section id='trend'><h2>趋势变化（距上次 {result['trend_days']} 天）</h2>
  <div class='card'><h3>增长最快的商品（评价增长 ≈ 销量）</h3>{_table(t, tc, title_col='商品名', own_col='是我的店', max_rows=30)}</div>
  <div class='card'><h3>价格变动</h3>{_table(price_moves, tc, title_col='商品名', own_col='是我的店', max_rows=30)}</div>
</section>""")
    else:
        sections.append("<section><div class='card muted'>还没有历史数据：每天/每周运行一次后，这里会显示竞品的评价增长（估算销量）、价格和排名变化。</div></section>")

    page = TEMPLATE.replace("{{DATE}}", date.today().isoformat()).replace("{{SHOPS}}", html.escape(shops)) \
        .replace("{{NAV}}", "".join(nav)).replace("{{SECTIONS}}", "".join(sections)) \
        .replace("{{CHARTS}}", json.dumps(charts, ensure_ascii=False))
    path.write_text(page, encoding="utf-8")


TEMPLATE = """<!doctype html><html lang="zh"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>乐天竞品分析 {{DATE}}</title>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<style>
:root{--bg:#f6f6f4;--card:#fff;--fg:#1d1d1b;--muted:#6b6b66;--line:#e4e3de;--red:#bf0000;--own:#fff4d6;--c1:#9a9a94;--c2:#bf0000;}
@media (prefers-color-scheme:dark){:root{--bg:#161615;--card:#21211f;--fg:#ecebe6;--muted:#9d9c95;--line:#353531;--own:#3b3218;--c1:#77766f;--c2:#ff5a4f;}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.6 -apple-system,"Hiragino Sans","PingFang SC","Microsoft YaHei",sans-serif}
header{background:var(--red);color:#fff;padding:20px 16px}header h1{margin:0;font-size:22px}header p{margin:4px 0 0;opacity:.85}
nav{position:sticky;top:0;background:var(--card);border-bottom:1px solid var(--line);padding:8px 16px;display:flex;gap:16px;overflow-x:auto;z-index:2}
nav a{color:var(--fg);text-decoration:none;white-space:nowrap;font-weight:600}nav a:hover{color:var(--red)}
main{max-width:1280px;margin:0 auto;padding:16px}section{margin-bottom:40px}h2{border-left:5px solid var(--red);padding-left:10px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:16px;margin:14px 0}.card h3{margin:0 0 10px;font-size:16px}
.kpis{display:grid;grid-template-columns:repeat(auto-fill,minmax(140px,1fr));gap:10px}
.kpi{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px}.kpi .k{color:var(--muted);font-size:12px}.kpi .v{font-size:20px;font-weight:700;font-variant-numeric:tabular-nums}
.ins li{margin:4px 0}.muted{color:var(--muted)}.tw{overflow-x:auto}
table{border-collapse:collapse;width:100%;font-variant-numeric:tabular-nums}th,td{border-bottom:1px solid var(--line);padding:6px 8px;text-align:right;white-space:nowrap;vertical-align:top}
th{color:var(--muted);font-weight:600;font-size:12px}td.name{text-align:left;white-space:normal;min-width:240px;max-width:460px}
th:first-child,td:first-child,td.txt{text-align:left}tr.own td{background:var(--own)}a{color:var(--red)}
.chart{height:260px;margin-bottom:12px}
</style></head><body>
<header><h1>乐天竞品分析报告</h1><p>{{DATE}} ・ 我的店铺：{{SHOPS}} ・ 黄色底色 = 我的店铺/商品</p></header>
<nav>{{NAV}}</nav><main>{{SECTIONS}}</main>
<script>
const charts={{CHARTS}};const cs=getComputedStyle(document.documentElement);
if(!window.Chart){document.querySelectorAll('.chart').forEach(e=>e.style.display='none');}else{charts.forEach(c=>{new Chart(document.getElementById(c.id),{type:'bar',
data:{labels:c.labels,datasets:[{label:'商品占比（供给）',data:c.items,backgroundColor:cs.getPropertyValue('--c1')},
{label:'评价占比（需求）',data:c.reviews,backgroundColor:cs.getPropertyValue('--c2')}]},
options:{maintainAspectRatio:false,plugins:{tooltip:{callbacks:{label:x=>x.dataset.label+': '+Math.round(x.raw*100)+'%',
afterBody:i=>'我的商品数: '+c.mine[i[0].dataIndex]}}},scales:{y:{ticks:{callback:v=>Math.round(v*100)+'%'}}}}});});}
</script></body></html>"""
